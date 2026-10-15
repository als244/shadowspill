#!/usr/bin/env python3
"""Turn a quickstart timeline HTML/JSON or report directory into slide assets.

No training run, browser, GPU, or Google credentials are needed. Run with
python -m benchmarking.quickstart_timeline; see quickstart_timeline.md.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.patches import Patch, Rectangle
from matplotlib.ticker import MaxNLocator, StrMethodFormatter

COLORS = dict(
    params="#3268DE",
    gradients="#008D80",
    optimizer_state="#8860C9",
    activations="#E1AD4F",
    forward="#D65A22",
    backward="#75452E",
    optimizer="#151F2B",
    recompute="#84CC16",
    idle="#FCE9E7",
    ink="#15263D",
    muted="#62738A",
    line="#D6DFEB",
    red="#CF5555",
    occupancy_track="#EEF2F7",
)
MEMORY_CLASSES = {
    "params": "Parameters",
    "gradients": "Gradients",
    "optimizer_state": "Optimizer state",
    "activations": "Activations",
}


def budget_label(budget):
    """A budget as the axis names it: an 8.0137 GiB execution budget is the 8 GiB
    point, as in the tradeoff exporter."""
    whole = round(budget)
    return f"{whole:g}" if abs(budget - whole) < 0.05 else f"{budget:.2f}"


def page_budget(data):
    """The execution budget a simulated or traced page says it was planned for."""
    match = re.search(r"execution budget ([\d.]+)gib", data.get("plan", ""))
    return float(match[1]) if match else None


def report_root(path):
    """The quickstart run a page or summary sits in: the directory above timelines/."""
    parts = Path(path).resolve().parts
    return Path(*parts[: parts.index("timelines")]) if "timelines" in parts else None


def requested_budgets(root):
    """The device budgets a quickstart was asked for, from its request.json."""
    path = root / "request.json" if root is not None else None
    if path is None or not path.is_file():
        return []
    request = json.loads(path.read_text()).get("request", {})
    return sorted(
        {
            float(budget)
            for key in ("search_budget_gib", "run_budget_gib")
            for budget in request.get(key) or ()
        }
    )


def requested_for(execution, requested):
    """The device budget an execution budget was cut from. A run takes its
    headroom off each request, so that is the smallest request at or above the
    execution budget, within 2 GiB; None when no request fits."""
    above = [r for r in requested if 0 <= r - execution < 2]
    return min(above) if above else None


def step_program(page):
    """The program a page's step was planned on, and its path, found the way the
    report's own timeline writer names it: the stored plans label each program with
    its geometry, which the page's directory carries. None when ShadowSpill is not
    importable, the page is not inside a report's timelines/, or the program is not
    stored."""
    try:
        from shadowspill.diagnostics.occupancy import program_path_for, stored_plans
    except ImportError:
        return None
    parts = page.resolve().parts
    if "timelines" not in parts:
        return None
    root = Path(*parts[: parts.index("timelines")])
    folder = page.parent
    geometry = (folder.parent if folder.name.startswith("recompute_") else folder).name
    stored = stored_plans(root)
    plans = [plan for plan in stored if plan.geometry == geometry]
    if not plans and len({plan.program_digest for plan in stored}) == 1:
        plans = stored  # one program: its folders are named by plan, not geometry
    digests = {plan.program_digest for plan in plans}
    if len(digests) != 1:
        return None
    path = program_path_for(plans[0].path, {"program_digest": digests.pop()})
    return (json.loads(path.read_text()), path) if path.is_file() else None


def count_line(data, planned, program, transfers):
    """How much the step moves through memory: the objects its tasks use, the
    tasks (control tasks are bookkeeping, not work), fetches and evicts. The
    objects come from ``planned``, the page whose tasks carry program task ids."""
    tasks = [item for item in data["compute"] if item[2] != "control"]
    counts = {"tasks": len(tasks)}
    if program is not None:
        by_id = {task["task_id"]: task for task in program["tasks"]}
        run = [by_id.get(item[3]) for item in planned["compute"]]
        if all(run):
            used = {  # a reference is an object id, or a record naming one
                ref["object_id"] if isinstance(ref, dict) else ref
                for task in run
                for key in ("inputs", "outputs", "mutations")
                for ref in task.get(key) or ()
            }
            groups = {o["object_id"]: o["alias_group_id"] for o in program["objects"]}
            counts = {
                "objects": len(used),
                "alias_groups": len({groups[o] for o in used if o in groups}),
                **counts,
            }
    if transfers:
        counts.update(fetches=len(data["fetch"]), evicts=len(data["evict"]))
    order = ("objects", "tasks", "fetches", "evicts")
    text = " · ".join(f"{counts[k]:,} {k}" for k in order if k in counts)
    return counts, text


def category(name):
    """Presentation convention: temporary workspace is included in activations."""
    return {
        "weights": "params",
        "model gradients": "gradients",
        "optimizer state": "optimizer_state",
    }.get(name, "activations")


def select_source(args):
    if args.source.is_file():
        return args.source.resolve()
    summaries = list(args.source.rglob("timelines/summary.csv"))
    direct = args.source / "summary.csv"
    if direct.is_file():
        summaries.append(direct)
    summaries = sorted(set(p.resolve() for p in summaries))
    if len(summaries) != 1:
        raise ValueError(
            "Choose a single report's seq*/seqsperstep*/ directory; "
            f"found {len(summaries)} timeline summaries: {summaries}"
        )
    with summaries[0].open() as f:
        rows = list(csv.DictReader(f))
    if args.list:
        for row in rows:
            print(
                " | ".join(
                    row.get(k, "")
                    for k in (
                        "budget",
                        "geometry",
                        "resolution",
                        "view",
                        "step_seconds",
                        "page",
                    )
                )
            )
        return None
    view = (args.view or "traced").replace("-", "_")
    # Choose the simulated winner, then locate its requested timeline. A
    # 'selected' row alone is only the best resolution within ONE geometry.
    selection_view = (
        "simulated" if view in ("simulated", "traced", "unconstrained") else view
    )
    candidates = [r for r in rows if r["view"] == selection_view]
    if selection_view == "simulated" and not args.resolution:
        candidates = [
            r for r in candidates if r.get("selected", "true").lower() in ("true", "1")
        ]
    if args.budget_gib is not None:
        # Named as on the axis: 8 selects the 8.0137 GiB execution budget, or with
        # --budget-display requested, the execution budget an 8 GiB request left.
        requested = requested_budgets(summaries[0].parent.parent)

        def shown(budget):
            if args.budget_display == "execution":
                return budget
            return requested_for(budget, requested)

        budgets = sorted(
            {
                budget
                for budget in (
                    float(r["budget"][:-3])
                    for r in candidates
                    if r["budget"].endswith("gib")
                )
                if shown(budget) is not None
                and abs(shown(budget) - args.budget_gib) < 0.05
            }
        )
        if len(budgets) > 1:
            raise ValueError(
                f"Budget {args.budget_gib:g} GiB does not identify one budget: "
                f"{budgets}"
            )
        candidates = [
            r
            for r in candidates
            if budgets
            and r["budget"].endswith("gib")
            and math.isclose(float(r["budget"][:-3]), budgets[0], abs_tol=1e-6)
        ]
    elif selection_view == "simulated":
        raise ValueError(
            "A report directory needs --budget-gib; or pass the exact timeline HTML."
        )
    if args.geometry:
        candidates = [r for r in candidates if r["geometry"] == args.geometry]
    if args.resolution:
        candidates = [r for r in candidates if r["resolution"] == args.resolution]
    if not candidates:
        raise ValueError("No matching timeline. Use --list to inspect the report.")
    if selection_view == "all_save" and len(candidates) > 1:
        raise ValueError(
            "Multiple all-save geometries: specify --geometry or the exact HTML."
        )
    winner = min(candidates, key=lambda r: float(r["step_seconds"]))
    page = summaries[0].parent / winner["page"]
    if view != selection_view:
        page = page.with_name(view + ".html")
    if not page.is_file():
        raise ValueError(f"The selected plan has no {view} timeline: {page}")
    return page


def load_timeline(path):
    content = path.read_text()
    if path.suffix.lower() == ".html":
        match = re.search(
            r"<script\b[^>]*\bid=[\"\']data[\"\'][^>]*>(.*?)</script>", content, re.S
        )
        if not match:
            raise ValueError(f"No embedded quickstart timeline JSON in {path}")
        content = match[1]
    data = json.loads(content)
    required = {"pools", "compute", "fetch", "evict", "end_seconds", "summary", "view"}
    if missing := required - data.keys():
        raise ValueError(f"Missing timeline fields: {sorted(missing)}")
    return data


def lane_runs(items, duration, *, compute=False, columns=1800, occupancy=False):
    """Same dominant-duration overview binning as the report's HTML timeline.

    With ``occupancy``, a transfer lane instead returns one run per column it
    touches, with the share of that column the lane is busy as a fourth field:
    serial lanes never overlap themselves, so the shares never exceed one."""
    tally = [{} for _ in range(columns)]

    def paint(start, end, kind):
        a, b = (
            max(0, start) / duration * columns,
            min(duration, end) / duration * columns,
        )
        for i in range(max(0, math.floor(a)), min(columns - 1, math.floor(b)) + 1):
            part = min(b, i + 1) - max(a, i)
            if part > 0:
                tally[i][kind] = tally[i].get(kind, 0) + part

    items = sorted(items, key=lambda item: item[0])
    for item in items:
        a, b = item[:2]
        if compute:
            split = min(b, a + item[4])
            paint(a, split, "recompute")
            paint(split, b, item[2])
        else:
            paint(a, b, category(item[3]))
    if compute and items:
        latest = items[0][1]
        for item in items[1:]:
            if item[0] > latest:
                paint(latest, item[0], "idle")
            latest = max(latest, item[1])
    kinds = [max(t, key=t.get) if t else None for t in tally]
    if occupancy and not compute:
        return [
            (
                i * duration / columns,
                (i + 1) * duration / columns,
                kinds[i],
                min(1.0, sum(t.values())),
            )
            for i, t in enumerate(tally)
            if t
        ]
    runs, start = [], 0
    for i in range(1, columns + 1):
        if i == columns or kinds[i] != kinds[start]:
            if kinds[start] is not None:
                runs.append(
                    (start * duration / columns, i * duration / columns, kinds[start])
                )
            start = i
    return runs


def render(data, source, args):
    pool = next(p for p in data["pools"] if p["name"] == "execution")
    duration = float(data["end_seconds"])
    xmax = args.time_max or duration
    if xmax < duration:
        raise ValueError("--time-max must include the full timeline.")
    # The line sits at the page's own budget, or at the device budget it was cut
    # from with --budget-display requested; --budget-gib may only name it.
    budget = page_budget(data) if data["view"] in ("simulated", "traced") else None
    shown = budget
    if budget is not None and args.budget_display == "requested":
        shown = requested_for(budget, requested_budgets(report_root(source)))
        if shown is None:
            raise ValueError(
                f"request.json names no requested budget for the {budget:g} GiB "
                "execution budget."
            )
    if budget is None:
        budget = shown = args.budget_gib
    elif args.budget_gib is not None and abs(shown - args.budget_gib) >= 0.05:
        raise ValueError(
            f"--budget-gib {args.budget_gib:g} does not name this page's "
            f"{shown:g} GiB budget."
        )
    if args.persistent_gradients_gib is not None and data["view"] not in (
        "all_save",
        "unconstrained",
    ):
        raise ValueError(
            "Persistent-gradient illustration is only allowed "
            "for unconstrained timelines."
        )
    show_transfers = args.transfers == "show" or (
        args.transfers == "auto" and bool(data["fetch"] or data["evict"])
    )
    found = step_program(source) if args.counts else None
    program, program_path = found if found else (None, None)
    planned = data
    sibling = source.with_name("simulated.html")
    if data["view"] == "traced" and sibling.is_file():
        # A trace names its tasks for the device; the plan it ran names them by id.
        planned = load_timeline(sibling)
    step_counts, count_text = count_line(data, planned, program, show_transfers)
    if args.transfers == "hide" and (data["fetch"] or data["evict"]):
        raise ValueError("Cannot hide transfers that exist in the source timeline.")
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 9.5,
            "text.color": COLORS["ink"],
            "axes.labelcolor": COLORS["muted"],
            "axes.edgecolor": COLORS["line"],
            "xtick.color": COLORS["muted"],
            "ytick.color": COLORS["muted"],
            "svg.fonttype": "none",
            "pdf.fonttype": 42,
        }
    )
    fig = plt.figure(figsize=(12, 3.65), facecolor="white")
    left, width = 0.063, 0.926
    ax_mem = fig.add_axes(
        [
            left,
            0.50 if show_transfers else 0.34,
            width,
            0.405 if show_transfers else 0.565,
        ]
    )
    lanes = [("compute", "GPU", 0.355 if show_transfers else 0.155)]
    if show_transfers:
        lanes.extend([("fetch", "Fetch", 0.245), ("evict", "Evict", 0.135)])
    times = np.asarray(pool["times"])
    rows = {k: np.zeros(len(times)) for k in MEMORY_CLASSES}
    for name, values in pool["rows"].items():
        values = np.asarray(values)
        if len(values) != len(times) or not np.isfinite(values).all():
            raise ValueError(f"Invalid memory samples for {name}")
        rows[category(name)] += values
    original_sample_peak = float(np.max(sum(rows.values())))
    if args.persistent_gradients_gib is not None:
        rows["gradients"][:] = args.persistent_gradients_gib
    peak = (
        float(np.max(sum(rows.values())))
        if args.persistent_gradients_gib is not None
        else float(pool["peak_gib"])
    )
    ax_mem.stackplot(
        times,
        *rows.values(),
        colors=[COLORS[k] for k in rows],
        step="post",
        linewidth=0,
    )
    ymax = max(peak, shown or 0) * 1.15
    if shown:
        ax_mem.axhline(shown, color=COLORS["red"], lw=0.9, dashes=(4, 3))
        ax_mem.text(
            xmax,
            shown + ymax * 0.012,
            f"{budget_label(shown)} GiB budget",
            color=COLORS["red"],
            fontsize=9,
            ha="right",
            va="bottom",
            bbox=dict(facecolor="white", edgecolor="none", pad=0.3),
        )
        top = float(budget_label(shown))
        ax_mem.set_yticks([0, top / 2, top])
        ax_mem.yaxis.set_major_formatter(StrMethodFormatter("{x:g}"))
    else:
        ax_mem.yaxis.set_major_locator(MaxNLocator(nbins=5, steps=[1, 2, 5, 10]))
    ax_mem.set_ylim(0, ymax)
    ax_mem.tick_params(axis="y", length=0, pad=5)
    ax_mem.tick_params(axis="x", bottom=False, labelbottom=False)
    fig.text(left, 0.965, "GPU memory (GiB)", fontsize=11, weight="bold", va="top")
    peak_label = (
        f"Illustrative peak ≈{peak:.0f} GiB"
        if args.persistent_gradients_gib is not None
        else f"Peak {peak:.1f} GiB"
    )
    fig.text(0.989, 0.965, peak_label, fontsize=11, weight="bold", va="top", ha="right")
    axes, counts, used, optimizer_marks = [ax_mem], {}, set(), 0
    for key, label, bottom in lanes:
        ax = fig.add_axes([left, bottom, width, 0.077])
        axes.append(ax)
        occupancy = key != "compute" and args.transfer_lanes == "occupancy"
        runs = lane_runs(
            data[key],
            duration,
            compute=key == "compute",
            columns=args.columns,
            occupancy=occupancy,
        )
        counts[key] = len(runs)
        if occupancy:  # each column as tall as the share of it the lane is busy
            ax.axhspan(0, 1, color=COLORS["occupancy_track"], lw=0, zorder=0)
            for kind in sorted({run[2] for run in runs}):
                mine = [run for run in runs if run[2] == kind]
                ax.bar(
                    [a for a, *_ in mine],
                    [share for *_, share in mine],
                    width=[b - a for a, b, *_ in mine],
                    align="edge",
                    color=COLORS[kind],
                    linewidth=0,
                )
            runs = []
        for a, b, kind in runs:
            ax.broken_barh(
                [(a, b - a)], (0, 1), facecolors=COLORS[kind], edgecolors="none"
            )
            if key == "compute":
                used.add(kind)
        if key == "compute":
            # An optimizer update lasts milliseconds and loses its overview column
            # to the backward work around it, so each one is marked where it runs:
            # a short tick just above the lane, which leaves the lane's work visible.
            updates = [item for item in data[key] if item[2] == "optimizer"]
            for a, b, *_ in updates:
                ax.axvline(
                    (a + b) / 2,
                    ymin=1.04,
                    ymax=1.3,
                    color=COLORS["optimizer"],
                    lw=0.9,
                    clip_on=False,
                    solid_capstyle="butt",
                )
            if updates:
                used.add("optimizer")
            optimizer_marks = len(updates)
            for a, b, *_ in data[key]:
                if (b - a) / xmax * args.columns > 3:
                    ax.add_patch(
                        Rectangle(
                            (a, 0),
                            b - a,
                            1,
                            facecolor="none",
                            edgecolor=(0, 0, 0, 0.20),
                            linewidth=0.25,
                        )
                    )
        ax.set_ylim(0, 1)
        ax.set_yticks([])
        ax.text(
            -0.018,
            0.5,
            label,
            transform=ax.transAxes,
            color=COLORS["muted"],
            fontsize=10,
            ha="right",
            va="center",
        )
        ax.tick_params(axis="x", length=0, pad=4, labelbottom=key == lanes[-1][0])
    ticks = MaxNLocator(nbins=10, steps=[1, 2, 2.5, 5, 10]).tick_values(0, xmax)
    ticks = ticks[(ticks >= 0) & (ticks <= xmax)]
    for ax in axes:
        ax.set_xlim(0, xmax)
        ax.set_xticks(ticks)
        for side in ("top", "right", "left"):
            ax.spines[side].set_visible(False)
        ax.spines["bottom"].set_linewidth(0.6)
    axes[-1].text(
        1,
        -0.9,
        "Time (s)",
        transform=axes[-1].transAxes,
        fontsize=9,
        ha="right",
        color=COLORS["muted"],
    )
    handles = [
        Patch(facecolor=COLORS[k], edgecolor="none", label=name)
        for k, name in (
            ("forward", "Forward"),
            ("backward", "Backward"),
            ("recompute", "Recompute"),
            ("optimizer", "Optimizer Update"),
            ("idle", "Stalled"),
        )
        if k in used
    ]
    fig.legend(
        handles=handles,
        loc="lower left",
        bbox_to_anchor=(left - 0.004, -0.007),
        ncol=5,
        frameon=False,
        fontsize=9,
        handlelength=1.1,
        columnspacing=1.8,
        borderaxespad=0,
    )
    if args.counts:  # the legend row's free right end
        fig.text(
            left + width,
            0.012,
            count_text,
            fontsize=9,
            ha="right",
            va="bottom",
            color=COLORS["muted"],
        )
    args.outdir.mkdir(parents=True, exist_ok=True)
    for ext in ("png", "svg", "pdf"):
        fig.savefig(
            args.outdir / f"timeline.{ext}",
            dpi=args.dpi,
            transparent=ext != "pdf",
            facecolor="white" if ext == "pdf" else "none",
        )
    plt.close(fig)
    meta = dict(
        source=str(source),
        source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        view=data["view"],
        clock=data.get("clock"),
        plan=data.get("plan"),
        step_seconds=data["summary"]["seconds"],
        duration_seconds=duration,
        tokens_per_second=data["summary"].get("tokens_per_second"),
        time_max=xmax,
        peak_gib=peak,
        original_peak_gib=pool["peak_gib"],
        original_sample_peak_gib=original_sample_peak,
        budget_gib=budget,
        budget_display=args.budget_display,
        shown_budget_gib=shown,
        pool_capacity_gib=pool.get("capacity_gib"),
        persistent_gradients_gib=args.persistent_gradients_gib,
        spill_peak_gib=data["summary"].get("spill_peak_gib"),
        events={k: len(data[k]) for k in ("compute", "fetch", "evict")},
        counts=dict(step_counts, program=str(program_path) if program_path else None),
        summary=data["summary"],
        rendered_runs=counts,
        transfer_lanes=args.transfer_lanes,
        optimizer_marks=optimizer_marks,
        overview_columns=args.columns,
        interpretation={
            "memory": (
                "Report tensor leases plus task workspace envelope; "
                "not all physical CUDA allocations."
            ),
            "categories": (
                "Workspace and control temporaries are included in Activations; "
                "no bytes dropped."
            ),
            "recompute": (
                "Profile-cost split within backward tasks, as in source HTML; "
                "not separately traced kernels."
            ),
            "idle": "Only between compute tasks; initial/final waiting is not shaded.",
            "transfers": (
                "Colors identify object class. Overlapping/narrow events "
                "use dominant-duration overview bins."
            ),
            "illustration": (
                "Persistent gradients replace the original series "
                "only when explicitly requested."
            ),
        },
    )
    (args.outdir / "metadata.json").write_text(json.dumps(meta, indent=2) + "\n")
    if args.pptx:
        make_slide(args, meta)
    return meta


def make_slide(args, meta):
    """Portable slide using the existing deck's result-slide design."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("F7F9FC")

    def text(value, x, y, width, height, size, color="ink", bold=False):
        box = slide.shapes.add_textbox(
            Inches(x), Inches(y), Inches(width), Inches(height)
        )
        tf = box.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = value
        p.font.name, p.font.size, p.font.bold = "Arial", Pt(size), bold
        p.font.color.rgb = RGBColor.from_string(COLORS[color][1:])

    def rect(x, y, width, height, color, rounded=False):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(width),
            Inches(height),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color.lstrip("#"))
        shape.line.fill.background()
        if rounded:
            shape.adjustments[0] = 0.08

    label = {
        "traced": "MEASURED",
        "simulated": "SIMULATED",
        "all_save": "BEFORE PLANNING",
    }.get(meta["view"], "TIMELINE")
    text(f"SHADOWSPILL / {label}", 0.62, 0.40, 12.1, 0.3, 11, "muted", True)
    text(
        args.title or f"{label.title()} training timeline",
        0.62,
        0.86,
        12.1,
        0.60,
        30,
        bold=True,
    )
    subtitle = args.subtitle or (
        f"{meta['view']} · {meta['step_seconds']:.2f} s · "
        f"peak {meta['peak_gib']:.2f} GiB"
    )
    show_summary = args.summary and meta["view"] in ("simulated", "traced")
    text(subtitle, 0.64, 1.50 if show_summary else 1.63, 12.05, 0.38, 15, "muted")
    for x, (key, label) in zip(
        (0.76, 2.32, 3.88, 5.86), MEMORY_CLASSES.items(), strict=True
    ):
        rect(x, 1.975 if show_summary else 2.185, 0.13, 0.13, COLORS[key])
        text(label, x + 0.2, 1.945 if show_summary else 2.155, 1.75, 0.3, 11, "muted")
    rect(0.62, 2.15 if show_summary else 2.46, 12.11, 3.75, "FFFFFF", True)
    slide.shapes.add_picture(
        str(args.outdir / "timeline.png"),
        Inches(0.80),
        Inches(2.21 if show_summary else 2.53),
        width=Inches(11.73),
    )
    if show_summary:
        rect(0.62, 5.95, 12.11, 1.12, "E9EFF8", True)
        for i, (label, value) in enumerate(summary_cards(meta["summary"])):
            x, y = 0.78 + (i % 6) * 1.995, 6.02 + (i // 6) * 0.52
            text(label, x, y, 1.89, 0.21, 12, "muted")
            text(value, x, y + 0.20, 1.89, 0.25, 14, bold=True)
    else:
        rect(0.62, 6.40, 12.11, 0.50, "E9EFF8", True)
        text(
            args.caption
            or f"{meta['step_seconds']:.2f} s per step · {meta['view']} timeline",
            0.82,
            6.49,
            11.7,
            0.30,
            15,
            bold=True,
        )
    slide.notes_slide.notes_text_frame.text = json.dumps(meta, indent=2)
    deck.save(args.outdir / "timeline-slide.pptx")


def summary_cards(summary):
    """The twelve summary statistics from the quickstart HTML, in slide order,
    rounded as a slide reads them: no more digits than the comparison needs."""

    def number(key, spec, unit=""):
        value = summary.get(key)
        return "n/a" if value is None else f"{value:{spec}}{unit}"

    def percent(key):
        value = summary.get(key)
        if value is None:
            return "n/a"
        return "0%" if value == 0 else f"{value:.{1 if value < 10 else 0}f}%"

    def rate(direction):
        value = number(direction + "_gbps", ".0f")
        assumed = number("assumed_" + direction + "_gbps", ".0f")
        return f"{value} / {assumed} GB/s"

    return [
        ("Step time", number("seconds", ".1f", " s")),
        ("Tokens/s", number("tokens_per_second", ",.0f")),
        ("Peak GPU memory", number("execution_peak_gib", ".1f", " GiB")),
        ("Peak spill", number("spill_peak_gib", ",.0f", " GiB")),
        ("Stalled", percent("idle_percent")),
        ("Recompute", percent("recompute_percent")),
        ("Fetch lane busy", percent("fetch_utilization_percent")),
        ("Evict lane busy", percent("evict_utilization_percent")),
        ("Fetched", number("fetch_gib", ",.0f", " GiB")),
        ("Evicted", number("evict_gib", ",.0f", " GiB")),
        ("Fetch rate / planned blend", rate("fetch")),
        ("Evict rate / planned blend", rate("evict")),
    ]


def main():
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.ArgumentDefaultsHelpFormatter
    )
    parser.add_argument(
        "source",
        type=Path,
        help=(
            "Timeline HTML/JSON, timelines directory, "
            "or one quickstart result directory"
        ),
    )
    parser.add_argument(
        "--view", choices=("traced", "simulated", "all-save", "unconstrained")
    )
    parser.add_argument(
        "--budget-gib",
        type=float,
        help="Select a budget's fastest simulated plan; also label the budget line",
    )
    parser.add_argument(
        "--budget-display",
        choices=("execution", "requested"),
        default="execution",
        help=(
            "Name and draw the budget as the execution budget left after headroom, "
            "or as the device budget request.json asked for"
        ),
    )
    parser.add_argument(
        "--geometry", help="Optional exact report geometry, e.g. 8x8_1x8rp"
    )
    parser.add_argument("--resolution", help="Optional resolution, e.g. 3/4")
    parser.add_argument(
        "--list",
        action="store_true",
        help="List a report directory's available timeline choices",
    )
    parser.add_argument(
        "--outdir",
        type=Path,
        default=Path("timeline-slide"),
        help="Directory for generated figures and slide",
    )
    parser.add_argument(
        "--time-max",
        type=float,
        help="Common x-axis endpoint for comparable plots; must cover all events",
    )
    parser.add_argument(
        "--transfers",
        choices=("auto", "show", "hide"),
        default="auto",
        help="Show Fetch/Evict lanes; auto omits empty lanes",
    )
    parser.add_argument(
        "--persistent-gradients-gib",
        type=float,
        help="Explicit illustration only: replace gradients with this constant bank",
    )
    parser.add_argument("--columns", type=int, default=1800, help="Overview time bins")
    parser.add_argument(
        "--transfer-lanes",
        choices=("occupancy", "presence"),
        default="occupancy",
        help=(
            "Fetch/evict columns: occupancy (default) draws each column as tall "
            "as the share of it the lane is busy; presence paints any column a "
            "transfer touches full height, the HTML overview's rule"
        ),
    )
    parser.add_argument("--dpi", type=int, default=300, help="PNG resolution")
    parser.add_argument(
        "--pptx",
        action="store_true",
        help="Also create a one-slide PowerPoint for import into Slides",
    )
    parser.add_argument(
        "--summary",
        action=argparse.BooleanOptionalAction,
        default=True,
        help=(
            "Include the HTML's summary statistics "
            "in simulated/traced PowerPoint slides"
        ),
    )
    parser.add_argument("--title", help="PowerPoint title")
    parser.add_argument("--subtitle", help="PowerPoint subtitle")
    parser.add_argument(
        "--counts",
        action=argparse.BooleanOptionalAction,
        default=True,
        help="Print the step's objects, tasks, fetches and evicts beside the legend",
    )
    parser.add_argument(
        "--caption", help="PowerPoint takeaway when summary cards are not shown"
    )
    args = parser.parse_args()
    try:
        if args.list and args.source.is_file():
            raise ValueError(
                "--list needs a report directory, not a single timeline file."
            )
        if args.columns <= 0 or args.dpi <= 0:
            raise ValueError("--columns and --dpi must be positive.")
        for name in ("budget_gib", "time_max", "persistent_gradients_gib"):
            value = getattr(args, name)
            if value is not None and (
                not math.isfinite(value)
                or value < 0
                or (name != "persistent_gradients_gib" and value == 0)
            ):
                raise ValueError(
                    f"--{name.replace('_', '-')} must be finite and positive."
                )
        source = select_source(args)
        if source is None:
            return
        data = load_timeline(source)
        if args.view and data["view"] != args.view.replace("-", "_"):
            raise ValueError(
                f"Source contains {data['view']!r}, not requested {args.view!r}."
            )
        meta = render(data, source, args)
    except (ValueError, OSError, KeyError, StopIteration) as exc:
        parser.error(str(exc))
    print(
        json.dumps(
            dict(
                outdir=str(args.outdir.resolve()),
                view=meta["view"],
                seconds=meta["step_seconds"],
                peak_gib=meta["peak_gib"],
            ),
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
