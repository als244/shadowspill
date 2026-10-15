"""Export a quickstart memory-budget/throughput comparison as a figure or slide.

Run ``python -m benchmarking.quickstart_tradeoff --help``. Input is an existing
report directory, its figures/ directory, or figures/raw_data/. No GPU is used.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import statistics
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.lines import Line2D
from matplotlib.path import Path as LinePath
from matplotlib.ticker import MaxNLocator, StrMethodFormatter

from benchmarking.quickstart_timeline import (
    COLORS,
    load_timeline,
    requested_budgets,
    requested_for,
)


def raw_data_directory(source: Path) -> Path:
    """A run's raw data: points.csv from its search, run_budgets.csv when it
    measured. A rerun that measured chosen plans without searching has only the
    second."""
    for path in (source, source / "raw_data", source / "figures/raw_data"):
        if (path / "points.csv").is_file() or (path / "run_budgets.csv").is_file():
            return path.resolve()
    raise ValueError(
        "Expected points.csv or run_budgets.csv in SOURCE, raw_data/, or "
        "figures/raw_data/. Select one quickstart run, not its parent collection."
    )


def planned_from_run(run: dict[str, str]) -> dict[str, str]:
    """A measured budget's plan as points.csv would record it, for a rerun with
    no search: the run row splits the same predicted step into task time (with
    its recompute), idle time, the entry delay and the final writeback."""
    metric = "units" if "simulated_units_per_second" in run else "tokens"
    recompute = float(run["recomputation_seconds"])
    idle = float(run["simulated_idle_seconds"]) + float(
        run["simulated_entry_delay_seconds"]
    )
    return {
        "status": "succeeded",
        "execution_budget_gib": run["execution_budget_gib"],
        "simulated_step_seconds": run["simulated_step_seconds"],
        f"simulated_{metric}_per_second": run[f"simulated_{metric}_per_second"],
        "unconstrained_seconds": str(float(run["profiled_task_seconds"]) - recompute),
        "recomputation_overhead_seconds": str(recompute),
        "idle_seconds": str(idle),
        "terminal_writeback_seconds": run["terminal_tail_seconds"],
    }


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="") as handle:
        return list(csv.DictReader(handle))


def positive(value: float, name: str) -> float:
    if not math.isfinite(value) or value <= 0:
        raise ValueError(f"{name} must be positive and finite, got {value}.")
    return value


def agree(actual: float, expected: float, name: str) -> None:
    if not math.isclose(actual, expected, rel_tol=1e-6, abs_tol=1e-7):
        raise ValueError(f"{name}: inconsistent report values {actual} and {expected}.")


def load_winners(
    directory: Path, budgets: list[float] | None, display: str = "execution"
):
    """Each budget's winning plan: measured where the run executed it, and only
    simulated for a report that planned without running (no run_budgets.csv).
    With ``display="requested"`` budgets are named, filtered and plotted by the
    device budget the run was asked for rather than the execution budget left
    after headroom."""
    runs_path = directory / "run_budgets.csv"
    runs = read_csv(runs_path) if runs_path.is_file() else []
    points_path = directory / "points.csv"
    points = (
        read_csv(points_path)
        if points_path.is_file()
        else [planned_from_run(run) for run in runs]
    )
    planned = [p for p in points if p["status"] == "succeeded"]
    if not runs and not planned:
        raise ValueError("The report has no planned budgets to plot.")
    # Current generic reports use units/s; archived text reports use tokens/s.
    columns = runs[0] if runs else planned[0]
    metric = (
        "units" if any(k.endswith("_units_per_second") for k in columns) else "tokens"
    )
    throughput_path = directory / "throughput.json"
    throughput = (
        json.loads(throughput_path.read_text()) if throughput_path.is_file() else {}
    )
    if "units_per_step" in throughput:
        units = float(throughput["units_per_step"])
    elif runs:
        units = float(runs[0][f"measured_{metric}_per_second"]) * float(
            runs[0]["measured_step_seconds"]
        )
    else:
        units = float(planned[0][f"simulated_{metric}_per_second"]) * float(
            planned[0]["simulated_step_seconds"]
        )
    units = positive(units, "units per step")
    label = throughput.get("unit_label", metric)
    requested = requested_budgets(directory.parent.parent)

    def shown(budget: float) -> float:
        if display == "execution":
            return budget
        named = requested_for(budget, requested)
        if named is None:
            raise ValueError(
                f"request.json names no requested budget for the {budget:g} GiB "
                "execution budget."
            )
        return named

    available = sorted({float(r["execution_budget_gib"]) for r in runs or planned})
    selected = set(available)
    if budgets:
        selected = set()
        for budget in budgets:
            matches = [v for v in available if abs(shown(v) - budget) < 0.05]
            if len(matches) != 1:
                raise ValueError(
                    f"Budget {budget:g} GiB does not identify one planned point; "
                    f"available: {[shown(v) for v in available]}"
                )
            selected.add(matches[0])
    steps_path = directory / "steps.csv"
    steps = read_csv(steps_path) if steps_path.is_file() else []
    by_budget = {float(r["execution_budget_gib"]): r for r in runs}
    if len(by_budget) != len(runs):
        raise ValueError("Duplicate run rows for one budget.")
    rows = []
    for budget in sorted(selected):
        positive(budget, "budget")
        run = by_budget.get(budget)
        candidates = [
            p
            for p in points
            if p["status"] == "succeeded"
            and math.isclose(float(p["execution_budget_gib"]), budget, abs_tol=1e-7)
        ]
        if not candidates:
            raise ValueError(
                f"No successful planning candidate for measured budget {budget:g}."
            )
        winner = min(candidates, key=lambda p: float(p["simulated_step_seconds"]))
        predicted = positive(float(winner["simulated_step_seconds"]), "predicted time")
        measured, samples = None, []
        if run is not None:
            agree(
                predicted,
                float(run["simulated_step_seconds"]),
                f"winner at {budget:g} GiB",
            )
            measured = positive(float(run["measured_step_seconds"]), "measured time")
            at_budget = sorted(
                (
                    s
                    for s in steps
                    if math.isclose(
                        float(s["execution_budget_gib"]), budget, abs_tol=1e-7
                    )
                ),
                key=lambda s: int(s.get("step") or 0),
            )
            samples = [float(s["seconds"]) for s in at_budget]
            if steps and not samples:
                raise ValueError(f"steps.csv has no measurements at {budget:g} GiB.")
            for sample in samples:
                positive(sample, "step time")
            # The runner's median leaves out its first step, which reconciles the
            # plan's initial state, and its last, which is traced; older reports
            # took every step. Whichever the report used names the sample count.
            middle = samples[1:-1] if len(samples) > 2 else samples
            if middle and math.isclose(
                measured, statistics.median(middle), rel_tol=1e-6, abs_tol=1e-7
            ):
                samples = middle
            elif samples:
                agree(
                    measured,
                    statistics.median(samples),
                    f"measured median at {budget:g} GiB",
                )
        effective = positive(
            float(winner["unconstrained_seconds"]), "effective compute"
        )
        recompute = float(winner["recomputation_overhead_seconds"])
        idle = float(winner["idle_seconds"]) + float(
            winner["terminal_writeback_seconds"]
        )
        if any(not math.isfinite(v) or v < 0 for v in (recompute, idle)):
            raise ValueError(f"Invalid overhead durations at {budget:g} GiB.")
        agree(
            effective + recompute + idle, predicted, f"time breakdown at {budget:g} GiB"
        )
        if run is not None:
            agree(
                units / measured,
                float(run[f"measured_{metric}_per_second"]),
                "measured throughput",
            )
            agree(
                units / predicted,
                float(run[f"simulated_{metric}_per_second"]),
                "simulated throughput",
            )
        rows.append(
            dict(
                budget_gib=budget,
                shown_gib=shown(budget),
                measured_seconds=measured,
                simulated_seconds=predicted,
                measured_rate=units / measured if measured else None,
                simulated_rate=units / predicted,
                chosen_compute_rate=units / (effective + recompute),
                effective_compute_seconds=effective,
                recompute_seconds=recompute,
                idle_seconds=idle,
                effective_compute_pct=100 * effective / predicted,
                recompute_pct=100 * recompute / predicted,
                idle_pct=100 * idle / predicted,
                samples=len(samples) if samples else None,
                candidate=winner.get(
                    "candidate", winner.get("sequences_per_microbatch", "")
                ),
                accumulation_count=winner.get("accumulation_count"),
                ordering=winner.get("ordering"),
            )
        )
    sources = [
        p for p in (points_path, runs_path, throughput_path, steps_path) if p.is_file()
    ]
    return rows, units, label, sources


def merged_winners(args):
    """Every report's winners, one row per budget. A later report replaces an
    earlier one at a budget unless only the earlier measured it, so a rerun of
    one budget fills in a point the first report could not run."""
    by_budget, units, unit_label, sources, directories = {}, None, None, [], []
    for source in args.source:
        directory = raw_data_directory(source)
        found, found_units, found_label, found_sources = load_winners(
            directory, args.budget_gib, args.budget_display
        )
        if units is not None:
            agree(found_units, units, "units per step across reports")
        units, unit_label = found_units, unit_label or found_label
        report = (
            directory.parent.parent
            if directory.parent.name == "figures"
            else directory.parent
        )
        for row in found:
            row["report"] = report.name
            kept = by_budget.get(row["shown_gib"])
            if (
                kept is None
                or row["measured_rate"] is not None
                or kept["measured_rate"] is None
            ):
                by_budget[row["shown_gib"]] = row
        sources += found_sources
        directories.append(directory)
    rows = [by_budget[budget] for budget in sorted(by_budget)]
    return rows, units, unit_label, sources, directories


def unconstrained_reference(directory: Path, units: float, args):
    source = args.unconstrained
    if source is None:
        root = (
            directory.parent.parent
            if directory.parent.name == "figures"
            else directory.parent
        )
        choices = list((root / "timelines/all_save").glob("*/all_save.html"))
        if choices:  # the fastest; among equally fast ones, the smallest peak

            def speed_then_peak(path):
                data = load_timeline(path)
                pool = next(p for p in data["pools"] if p["name"] == "execution")
                return float(data["summary"]["seconds"]), float(pool["peak_gib"])

            source = min(choices, key=speed_then_peak)
    if source is None:
        if args.unconstrained_peak_gib is not None:
            raise ValueError(
                "--unconstrained-peak-gib needs an unconstrained timeline."
            )
        return None
    source = source.resolve()
    data = load_timeline(source)
    if data["view"] not in ("all_save", "unconstrained"):
        raise ValueError(
            "--unconstrained must name an all-save/unconstrained timeline."
        )
    seconds = positive(float(data["summary"]["seconds"]), "unconstrained time")
    source_rate = data["summary"].get(
        "units_per_second", data["summary"].get("tokens_per_second")
    )
    if source_rate is not None:
        agree(units, float(source_rate) * seconds, "unconstrained units per step")
    pool = next(p for p in data["pools"] if p["name"] == "execution")
    peak = positive(float(pool["peak_gib"]), "unconstrained peak")
    return dict(
        source=str(source),
        seconds=seconds,
        rate=units / seconds,
        source_peak_gib=peak,
        display_peak_gib=args.unconstrained_peak_gib or peak,
        peak_overridden=args.unconstrained_peak_gib is not None,
    )


def range_ceiling(value: float) -> float:
    scale = 10 ** math.floor(math.log10(positive(value, "axis range")))
    return math.ceil(value / scale * 10) * scale / 10


def place_labels(fig, labels):
    """Put each endpoint label where it covers no plotted line, no other axes and no
    earlier label, inside the figure and never below or left of its own axes (the
    tick labels live there), and keeps the order of the values it is drawn beside:
    the larger of two labels at one end is the higher. It tries its preferred side
    of the point, then the other side, then further out, then beside the point. The
    curves can cross, so no fixed side per series keeps every label legible."""
    renderer = fig.canvas.get_renderer()
    page = fig.bbox
    placed = []  # (box, end, value)

    def put(ax, text, xy, style, dx, dy, ha, va):
        label = ax.annotate(
            text,
            xy,
            xytext=(dx, dy),
            textcoords="offset points",
            ha=ha,
            va=va,
            **style,
        )
        return label, label.get_window_extent(renderer).padded(2)

    def ordered(box, end, value):
        mid = (box.y0 + box.y1) / 2
        return all(
            (value - v) * (mid - (b.y0 + b.y1) / 2) > 0
            for b, e, v in placed
            if e == end and v != value
        )

    for ax, text, xy, above, align, style, end, value in labels:
        inside = ax.get_window_extent(renderer)
        others = [o.get_window_extent(renderer) for o in fig.axes if o is not ax]
        segments = []
        for line in ax.get_lines():
            points = line.get_transform().transform(
                np.column_stack([line.get_xdata(), line.get_ydata()])
            )
            segments += [LinePath(points[k : k + 2]) for k in range(len(points) - 1)]
        sides = (1, -1) if above else (-1, 1)
        flip = "right" if align == "left" else "left"
        inward = 1 if align == "left" else -1
        near = "bottom" if sides[0] > 0 else "top"
        candidates = [
            (0, side * distance, ha, "bottom" if side > 0 else "top")
            for ha in (align, flip)
            for distance in (7, 15)
            for side in sides
        ] + [
            (sx * 8, dy, "left" if sx > 0 else "right", va)
            for sx in (inward, -inward)
            for dy, va in ((0, "center"), (sides[0] * 4, near), (2, "bottom"))
        ]
        for candidate in candidates:
            label, box = put(ax, text, xy, style, *candidate)
            if (
                box.y0 >= inside.y0
                and box.x0 >= inside.x0 - 2
                and box.x1 <= page.x1 - 2
                and not any(box.overlaps(o) for o in others)
                and not any(box.overlaps(b) for b, _, _ in placed)
                and not any(s.intersects_bbox(box, filled=False) for s in segments)
                and ordered(box, end, value)
            ):
                break
            label.remove()
        else:  # nowhere is clear: the preferred spot
            label, box = put(ax, text, xy, style, *candidates[0])
        placed.append((box, end, value))


def render_figure(rows, reference, unit_label, args):
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 9.5,
            "text.color": COLORS["ink"],
            "axes.labelcolor": COLORS["muted"],
            "axes.edgecolor": COLORS["line"],
            "xtick.color": COLORS["muted"],
            "ytick.color": COLORS["muted"],
            "svg.fonttype": "none",
            "pdf.fonttype": 42,
        }
    )
    fig = plt.figure(figsize=(12, 4), facecolor="white")
    full = fig.add_axes([0.070, 0.575, 0.388, 0.245])
    zoom = fig.add_axes([0.070, 0.235, 0.388, 0.245], sharex=full)
    x = [r["shown_gib"] for r in rows]
    measured = all(r["measured_rate"] is not None for r in rows)
    series = ("measured_rate",) * measured + ("simulated_rate", "chosen_compute_rate")
    rates = [r[k] for r in rows for k in series]
    maximum = max(rates + ([reference["rate"]] if reference else []))
    full_max = args.full_ymax or range_ceiling(maximum * 1.06)
    if full_max <= maximum:
        raise ValueError("--full-ymax must be above all throughputs and the reference.")
    span = max(rates) - min(rates)
    padding = max(span * 0.30, max(rates) * 0.02)
    zoom_limits = args.zoom_ylim or [max(0, min(rates) - padding), max(rates) + padding]
    if not zoom_limits[0] < zoom_limits[1] or zoom_limits[0] < 0:
        raise ValueError("--zoom-ylim needs nonnegative increasing limits.")
    if zoom_limits[0] > min(rates) or zoom_limits[1] < max(rates):
        raise ValueError("--zoom-ylim must include every budget-dependent series.")
    for ax in (full, zoom):
        for key, color, style, label, marker in (
            ("measured_rate", "ink", "-", "Measured", "o"),
            ("simulated_rate", "muted", "--", "Simulated", "o"),
            (
                "chosen_compute_rate",
                "recompute",
                (0, (5, 3)),
                "Chosen recompute (no stalls)",
                None,
            ),
        ):
            if key not in series:
                continue
            ax.plot(
                x,
                [r[key] for r in rows],
                color=COLORS[color],
                ls=style,
                lw=1.6,
                marker=marker,
                markersize=3.5 if key == "measured_rate" else 3,
                label=label,
                zorder=4 if key == "measured_rate" else 3,
            )
        ax.yaxis.set_major_locator(MaxNLocator(nbins=3, steps=[1, 2, 2.5, 5, 10]))
        ax.yaxis.set_major_formatter(
            StrMethodFormatter("{x:,.0f}" if maximum >= 100 else "{x:g}")
        )
    full.set_ylim(0, full_max)
    zoom.set_ylim(*zoom_limits)
    if reference:
        full.axhline(reference["rate"], color=COLORS["forward"], lw=1.3, ls=(0, (5, 3)))
        peak = reference["display_peak_gib"]
        peak_text = f"{peak:.0f}" if peak >= 10 else f"{peak:.2g}"
        rate_text = (
            f"{reference['rate']:,.0f}"
            if maximum >= 100
            else f"{reference['rate']:.3g}"
        )
        full.annotate(
            f"Unconstrained ({peak_text} GiB): {rate_text}",
            (x[0], reference["rate"]),
            xytext=(0, 5),
            textcoords="offset points",
            ha="left",
            va="bottom",
            color=COLORS["forward"],
            fontsize=9,
            weight="bold",
        )
    # A divider separates the two scales without adding a text label.
    fig.add_artist(
        Line2D(
            [0.070, 0.458],
            [0.535, 0.535],
            transform=fig.transFigure,
            color=COLORS["ink"],
            lw=2.4,
            ls=(0, (0.4, 3.4)),
            dash_capstyle="round",
        )
    )
    labels = []  # endpoint labels, placed once every axis has its final limits
    ends = ((0, "left"), (-1, "right")) if len(rows) > 1 else ((0, "left"),)
    for key, rival, color in (
        ("measured_rate", "simulated_rate", "ink"),
        ("simulated_rate", "measured_rate", "muted"),
    ):
        if key not in series:
            continue
        for i, align in ends:
            value = rows[i][key]
            labels.append(
                (
                    zoom,
                    f"{value:,.0f}" if maximum >= 100 else f"{value:.3g}",
                    (x[i], value),
                    rows[i][rival] is None
                    or value > rows[i][rival]
                    or (value == rows[i][rival] and key == "measured_rate"),
                    align,
                    dict(
                        color=COLORS[color],
                        fontsize=9,
                        weight="bold",
                        bbox=dict(
                            facecolor="white", edgecolor="none", pad=0.8, alpha=0.88
                        ),
                    ),
                    ("rate", i),
                    value,
                )
            )
    fig.text(
        0.070,
        0.985,
        "Throughput: measured vs. simulated" if measured else "Throughput: simulated",
        fontsize=11.5,
        weight="bold",
        va="top",
    )
    counts = {r["samples"] for r in rows}
    count = next(iter(counts)) if len(counts) == 1 else None
    sampling = (
        "planned and simulated, not run"
        if not measured
        else f"measured median of {count} steps"
        if count
        else "recorded per-budget measurements"
    )
    fig.text(
        0.070,
        0.90,
        f"{unit_label.capitalize()}/s · {sampling}",
        color=COLORS["muted"],
        fontsize=9.5,
        va="top",
    )

    low_max = max(r[k] for r in rows for k in ("recompute_pct", "idle_pct"))
    high_min = min(r["effective_compute_pct"] for r in rows)
    cut_low = math.ceil(low_max / 5) * 5 + 5
    cut_high = math.floor(high_min / 5) * 5 - 5
    broken = args.share_axis == "auto" and cut_high - cut_low >= 20
    if broken:
        share_top = fig.add_axes([0.574, 0.615, 0.405, 0.195])
        share_low = fig.add_axes([0.574, 0.235, 0.405, 0.290], sharex=share_top)
        share_top.set_ylim(
            cut_high,
            min(
                100,
                math.ceil(max(r["effective_compute_pct"] for r in rows) / 5) * 5
                + (5 if max(r["effective_compute_pct"] for r in rows) % 5 < 1 else 0),
            ),
        )
        share_low.set_ylim(0, cut_low)
        for ax, edge in ((share_top, 0), (share_low, 1)):
            for side in (0, 1):
                ax.plot(
                    [side - 0.009, side + 0.009],
                    [edge - 0.022, edge + 0.022],
                    transform=ax.transAxes,
                    color=COLORS["muted"],
                    lw=0.8,
                    clip_on=False,
                )
        fig.text(
            0.979,
            0.563,
            f"{cut_low:g}–{cut_high:g}% omitted",  # noqa: RUF001 -- range typography
            color=COLORS["muted"],
            fontsize=8.5,
            va="center",
            ha="right",
        )
        share_axes = [share_top, share_low]
    else:
        share_top = share_low = fig.add_axes([0.574, 0.235, 0.405, 0.575])
        share_low.set_ylim(0, 105)
        share_axes = [share_low]
    handles = []
    for key, color, label, style in (
        ("effective_compute_pct", "backward", "Effective compute", "-"),
        ("recompute_pct", "recompute", "Recompute", "-"),
        ("idle_pct", "red", "Stalled", ":"),
    ):
        ax = share_top if key == "effective_compute_pct" else share_low
        values = [r[key] for r in rows]
        ax.plot(
            x,
            values,
            color=COLORS[color],
            ls=style,
            lw=1.75,
            marker="o",
            markersize=4.5,
        )
        # Of the two small shares, the larger is labeled above its point.
        rival = {"recompute_pct": "idle_pct", "idle_pct": "recompute_pct"}.get(key)
        for i, align in ends:
            above = (
                rival is None
                or values[i] > rows[i][rival]
                or (values[i] == rows[i][rival] and key == "recompute_pct")
            )
            labels.append(
                (
                    ax,
                    "0%" if values[i] == 0 else f"{values[i]:.1f}%",
                    (x[i], values[i]),
                    above,
                    align,
                    dict(
                        color=COLORS[color],
                        fontsize=9,
                        weight="bold",
                        bbox=dict(
                            facecolor="white", edgecolor="none", pad=0.6, alpha=0.88
                        ),
                    ),
                    ("share" if rival else key, i),
                    values[i],
                )
            )
        handles.append(
            Line2D(
                [],
                [],
                color=COLORS[color],
                ls=style,
                lw=1.75,
                marker="o",
                markersize=4,
                label=label,
            )
        )
    for ax in share_axes:
        if broken:
            lo, hi = ax.get_ylim()
            ax.set_yticks(list(range(int(lo), int(hi) + 1, 5)))
        ax.yaxis.set_major_formatter(StrMethodFormatter("{x:g}%"))
    fig.text(
        0.574,
        0.985,
        "Winning plans: where the step goes",
        fontsize=11.5,
        weight="bold",
        va="top",
    )
    fig.text(
        0.574,
        0.90,
        "Simulated step share · stalled includes final writeback",
        color=COLORS["muted"],
        fontsize=9.5,
        va="top",
    )

    margin = max((max(x) - min(x)) * 0.04, 0.5)
    for ax in [full, zoom, *share_axes]:
        ax.set_xlim(min(x) - margin, max(x) + margin)
        if len(x) <= 12:
            ax.set_xticks(
                x, [f"{v:.1f}".removesuffix(".0") for v in x], rotation=45, ha="right"
            )
        else:
            ax.xaxis.set_major_locator(MaxNLocator(nbins=9))
        ax.set_xlabel("GPU memory budget (GiB)", fontsize=10, labelpad=7)
        ax.grid(color="#BBC7D7", lw=0.85, zorder=0)
        for side in ("top", "right", "left"):
            ax.spines[side].set_visible(False)
        ax.spines["bottom"].set_linewidth(0.6)
        ax.tick_params(axis="both", length=0, pad=4)
    for ax in [full] + ([share_top] if broken else []):
        ax.tick_params(axis="x", labelbottom=False)
        ax.set_xlabel("")
    if broken:
        share_top.spines["bottom"].set_visible(False)
    place_labels(fig, labels)
    fig.legend(
        *full.get_legend_handles_labels(),
        loc="lower left",
        bbox_to_anchor=(0.062, -0.005),
        ncol=3,
        frameon=False,
        fontsize=9,
        handlelength=1.7,
        columnspacing=1,
        handletextpad=0.5,
        borderaxespad=0,
    )
    fig.legend(
        handles=handles,
        loc="lower left",
        bbox_to_anchor=(0.567, -0.005),
        ncol=3,
        frameon=False,
        fontsize=9,
        handlelength=1.5,
        columnspacing=1.15,
        handletextpad=0.5,
        borderaxespad=0,
    )
    args.outdir.mkdir(parents=True, exist_ok=True)
    for extension in ("png", "svg", "pdf"):
        fig.savefig(
            args.outdir / f"tradeoff.{extension}",
            dpi=args.dpi,
            transparent=extension != "pdf",
            facecolor="white" if extension == "pdf" else "none",
        )
    plt.close(fig)
    return dict(
        full_ylim=[0, full_max],
        zoom_ylim=list(zoom_limits),
        share_break=[cut_low, cut_high] if broken else None,
    )


def export_slide(args, metadata):
    """Use the existing result-slide composition, with native editable labels."""
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise ValueError(
            "--pptx requires python-pptx: python -m pip install python-pptx"
        ) from exc
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("F7F9FC")

    def text(value, x, y, width, height, size, *, bold=False, muted=False):
        box = slide.shapes.add_textbox(
            Inches(x), Inches(y), Inches(width), Inches(height)
        )
        tf = box.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = value
        p.font.name, p.font.size, p.font.bold = "Arial", Pt(size), bold
        p.font.color.rgb = RGBColor.from_string(COLORS["muted" if muted else "ink"][1:])

    def card(y, height, color):
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.62),
            Inches(y),
            Inches(12.11),
            Inches(height),
        )
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor.from_string(color)
        box.line.fill.background()
        box.adjustments[0] = 0.08

    rows = metadata["rows"]
    text(
        "SHADOWSPILL / MEMORY–THROUGHPUT TRADEOFF",  # noqa: RUF001 -- slide typography
        0.62,
        0.40,
        12.1,
        0.3,
        11,
        bold=True,
        muted=True,
    )
    text(args.title, 0.62, 0.86, 12.1, 0.60, 30, bold=True)
    subtitle = args.subtitle or (
        f"{metadata['units_per_step']:g} {metadata['unit_label']} per step · "
        f"{len(rows)} {'measured' if rows[0]['measured_rate'] else 'planned'} budgets"
    )
    text(subtitle, 0.64, 1.63, 12.05, 0.38, 15, muted=True)
    card(2.02, 4.19, "FFFFFF")
    slide.shapes.add_picture(
        str(args.outdir / "tradeoff.png"),
        Inches(0.80),
        Inches(2.13),
        width=Inches(11.73),
    )
    kind = "measured" if rows[0]["measured_rate"] else "simulated"
    first, best = rows[0], max(r[f"{kind}_rate"] for r in rows)
    caption = args.caption or (
        f"{first['shown_gib']:g} GiB achieves "
        f"{100 * first[f'{kind}_rate'] / best:.0f}% "
        f"of the best {kind} throughput."
    )
    card(6.40, 0.50, "E9EFF8")
    text(caption, 0.82, 6.49, 11.7, 0.30, 15, bold=True)
    slide.notes_slide.notes_text_frame.text = json.dumps(metadata, indent=2)
    deck.save(args.outdir / "tradeoff-slide.pptx")


def comma_floats(value: str) -> list[float]:
    try:
        result = [positive(float(v), "budget") for v in value.split(",")]
    except ValueError as exc:
        raise argparse.ArgumentTypeError(str(exc)) from exc
    return result


def main():
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.ArgumentDefaultsHelpFormatter
    )
    parser.add_argument(
        "source",
        type=Path,
        nargs="+",
        help=(
            "Quickstart run root, figures/, or figures/raw_data/. Several reports "
            "of one workload merge by budget: each budget comes from the last "
            "report that measured it, or else the last that planned it"
        ),
    )
    parser.add_argument(
        "--outdir",
        type=Path,
        default=Path("tradeoff-slide"),
        help="Generated figure/slide directory",
    )
    parser.add_argument(
        "--budget-gib",
        type=comma_floats,
        help="Optional comma-separated subset, e.g. 8,16,29",
    )
    parser.add_argument(
        "--budget-display",
        choices=("execution", "requested"),
        default="execution",
        help=(
            "Name budgets by the execution budget left after headroom, or by the "
            "device budget request.json asked for"
        ),
    )
    parser.add_argument(
        "--unconstrained",
        type=Path,
        help=(
            "All-save/unconstrained HTML or JSON; "
            "otherwise discover the fastest all-save timeline"
        ),
    )
    parser.add_argument(
        "--unconstrained-peak-gib",
        type=float,
        help=(
            "Explicit memory-label override for an adjusted illustration; "
            "source peak remains in metadata"
        ),
    )
    parser.add_argument(
        "--full-ymax",
        type=float,
        help="Full-range throughput upper bound; auto extends above every reference",
    )
    parser.add_argument(
        "--zoom-ylim",
        type=float,
        nargs=2,
        metavar=("MIN", "MAX"),
        help="Detail throughput range; auto fits all three budget-dependent lines",
    )
    parser.add_argument(
        "--share-axis",
        choices=("auto", "full"),
        default="auto",
        help=(
            "Break a large unused percentage range, or show the whole percentage scale"
        ),
    )
    parser.add_argument(
        "--unit-label",
        help="Throughput unit label; inferred from throughput.json or CSV columns",
    )
    parser.add_argument("--dpi", type=int, default=300, help="PNG resolution")
    parser.add_argument(
        "--pptx", action="store_true", help="Also export a one-slide PowerPoint"
    )
    parser.add_argument(
        "--title",
        default="Memory budget versus training throughput.",
        help="Editable slide title",
    )
    parser.add_argument("--subtitle", help="Editable slide subtitle")
    parser.add_argument(
        "--caption",
        help=(
            "Editable slide takeaway; default reports "
            "smallest-budget throughput relative to best"
        ),
    )
    args = parser.parse_args()
    try:
        positive(args.dpi, "DPI")
        for name in ("full_ymax", "unconstrained_peak_gib"):
            if getattr(args, name) is not None:
                positive(getattr(args, name), name)
        if args.zoom_ylim and not all(math.isfinite(v) for v in args.zoom_ylim):
            raise ValueError("--zoom-ylim must be finite.")
        rows, units, unit_label, sources, directories = merged_winners(args)
        reference = None
        for directory in directories if args.unconstrained is None else directories[:1]:
            found = unconstrained_reference(directory, units, args)
            if found and (
                reference is None
                or (found["rate"], -found["source_peak_gib"])
                > (reference["rate"], -reference["source_peak_gib"])
            ):
                reference = found  # the fastest all-save timeline of any report
        if reference:
            sources.append(Path(reference["source"]))
        axes = render_figure(rows, reference, args.unit_label or unit_label, args)
        metadata = dict(
            sources=[
                dict(path=str(p), sha256=hashlib.sha256(p.read_bytes()).hexdigest())
                for p in sources
            ],
            units_per_step=units,
            unit_label=args.unit_label or unit_label,
            axes=axes,
            unconstrained=reference,
            rows=rows,
            interpretation={
                "measured": (
                    "Recorded per-budget medians, "
                    "checked against steps.csv when present."
                ),
                "shares": (
                    "Simulated effective compute + added recompute + "
                    "idle/final-writeback time, divided by predicted step time."
                ),
                "chosen_compute": (
                    "Units per step / (effective compute + recompute); "
                    "a profile-cost ceiling without transfer stalls."
                ),
                "unconstrained": (
                    "All-save/unconstrained profile-cost reference; "
                    "not measured fully resident throughput."
                ),
                "divider": (
                    "Dark dotted line separates full-range and detail "
                    "throughput axes; it is not a data series."
                ),
            },
        )
        (args.outdir / "metadata.json").write_text(
            json.dumps(metadata, indent=2) + "\n"
        )
        with (args.outdir / "winners.csv").open("w", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
        if args.pptx:
            export_slide(args, metadata)
    except (ValueError, OSError, KeyError, StopIteration) as exc:
        parser.error(str(exc))
    print(
        json.dumps(
            dict(
                outdir=str(args.outdir.resolve()),
                budgets=len(rows),
                unconstrained_reference=reference is not None,
                pptx=args.pptx,
            ),
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
